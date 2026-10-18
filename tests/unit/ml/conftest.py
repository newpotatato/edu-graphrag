"""Tiny generated study archive shared by the ingestion and EDA tests."""

import importlib
import json
import shutil
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

PARAGRAPH = (
    "Topological space is a set with a family of open subsets closed under unions "
    "and finite intersections. "
)


def make_pdf(path: Path, text: str) -> None:
    # Imported dynamically: mypy crashes while analysing pymupdf's own sources.
    pymupdf = importlib.import_module("pymupdf")
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(pymupdf.Rect(40, 40, 560, 800), text, fontsize=10)
    pdf.save(str(path))
    pdf.close()


def make_ipynb(path: Path) -> None:
    """Minimal nbformat 4 notebook, written as JSON (nbformat's builders are untyped)."""
    notebook = {
        "nbformat": 4,
        "nbformat_minor": 5,
        "metadata": {
            "kernelspec": {"name": "python3", "display_name": "Python 3", "language": "python"}
        },
        "cells": [
            {
                "id": "c1",
                "cell_type": "markdown",
                "metadata": {},
                "source": ["# Lab 1\n", "\n", PARAGRAPH],
            },
            {
                "id": "c2",
                "cell_type": "code",
                "metadata": {},
                "execution_count": 1,
                "outputs": [{"output_type": "stream", "name": "stdout", "text": ["OUTPUT\n"]}],
                "source": "# not a heading\nx = 1",
            },
            {"id": "c3", "cell_type": "markdown", "metadata": {}, "source": ""},
        ],
    }
    path.write_text(json.dumps(notebook), encoding="utf-8")


def make_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Report", level=0)
    document.add_heading("Section", level=1)
    document.add_paragraph(PARAGRAPH)
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "metric"
    table.cell(0, 1).text = "value"
    document.save(str(path))


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Manifolds"
    slide.placeholders[1].text = PARAGRAPH
    rows = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = rows.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    shape.table.cell(0, 0).text = "a"
    shape.table.cell(1, 1).text = "b"
    presentation.save(str(path))


@pytest.fixture
def archive(tmp_path: Path) -> Path:
    """`<root>/` with sorted, unsorted and excluded files of every supported type."""
    root = tmp_path / "university"
    sorted_dir = root / "2_grade" / "4_sem" / "func_analyse"
    unsorted_dir = root / "3_grade" / "difgem"
    for directory in (
        sorted_dir,
        unsorted_dir,
        root / "3_grade" / "MLops" / "repo",
        root / "1_grade" / "1_sem" / "python" / "node_modules",
        root / "1_grade" / "1_sem" / "python" / ".hidden",
    ):
        directory.mkdir(parents=True)
    (sorted_dir / "notes.md").write_text(f"# Banach\n\n{PARAGRAPH * 3}", encoding="utf-8")
    (sorted_dir / "hw.tex").write_text(f"\\section{{Hilbert}}\n{PARAGRAPH}", encoding="utf-8")
    (sorted_dir / "image.png").write_bytes(b"not a document")
    make_ipynb(sorted_dir / "lab.ipynb")
    make_docx(sorted_dir / "report.docx")
    make_pptx(unsorted_dir / "slides.pptx")
    make_pdf(unsorted_dir / "lecture.pdf", PARAGRAPH * 4)
    shutil.copyfile(unsorted_dir / "lecture.pdf", unsorted_dir / "lecture (2).pdf")
    make_pdf(unsorted_dir / "scan.pdf", "")
    (unsorted_dir / "broken.docx").write_bytes(b"not a zip file")
    (root / "3_grade" / "MLops" / "repo" / "README.md").write_text(PARAGRAPH, encoding="utf-8")
    (root / "1_grade" / "1_sem" / "python" / "node_modules" / "x.md").write_text("x")
    (root / "1_grade" / "1_sem" / "python" / ".hidden" / "x.md").write_text("x")
    (root / "1_grade" / "1_sem" / "python" / ".~lock.md").write_text("x")
    return root
