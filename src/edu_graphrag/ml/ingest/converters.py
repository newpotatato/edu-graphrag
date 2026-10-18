"""Converters: raw document -> markdown / plain text.

Each converter takes a local file path and returns text. `convert` wraps the registry so a
broken file never stops the pipeline: the error is logged and returned in the result.
"""

import importlib
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

import structlog

from edu_graphrag.ml.ingest.connectors import RawDocument, SourceType

logger = structlog.get_logger(__name__)

Converter = Callable[[Path], str]
ConversionStatus = Literal["ok", "failed"]


def convert_pdf(path: Path, *, use_layout: bool = False) -> str:
    """Markdown via pymupdf4llm; falls back to plain pymupdf text if that fails.

    `use_layout` enables the pymupdf-layout GNN: better reading order, but much slower and
    memory hungry, so the classic heuristic parser is the default.
    """
    # Imported dynamically: mypy crashes while analysing pymupdf's own sources.
    pymupdf = importlib.import_module("pymupdf")
    pymupdf4llm = importlib.import_module("pymupdf4llm")

    try:
        pymupdf4llm.use_layout(use_layout)
        options = {"use_ocr": False} if use_layout else {}
        text: str = pymupdf4llm.to_markdown(str(path), show_progress=False, **options)
    except Exception as exc:
        logger.warning("pdf_markdown_failed_fallback", path=str(path), error=repr(exc))
        with pymupdf.open(str(path)) as pdf:
            text = "\n\n".join(page.get_text() for page in pdf)
    return text


def _cell_source(cell: Any) -> str:
    source = cell.get("source", "")
    return "".join(source) if isinstance(source, list) else str(source)


def convert_ipynb(path: Path) -> str:
    """Markdown cells as-is, code cells fenced; outputs are skipped."""
    import nbformat

    notebook = nbformat.read(str(path), as_version=4)  # type: ignore[no-untyped-call]
    language = notebook.metadata.get("kernelspec", {}).get("language") or "python"
    parts: list[str] = []
    for cell in notebook.cells:
        source = _cell_source(cell).strip()
        if not source:
            continue
        if cell.cell_type == "code":
            parts.append(f"```{language}\n{source}\n```")
        else:
            parts.append(source)
    return "\n\n".join(parts)


def convert_text(path: Path) -> str:
    """Markdown and LaTeX sources are already text."""
    return path.read_text(encoding="utf-8", errors="replace")


def convert_docx(path: Path) -> str:
    """Paragraphs (headings become markdown headings) followed by table rows."""
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and style[-1:].isdigit():
            text = "#" * int(style[-1]) + " " + text
        elif style == "Title":
            text = "# " + text
        parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def convert_pptx(path: Path) -> str:
    """One `## Slide N` section per slide with all text frames and tables."""
    import pptx

    presentation = pptx.Presentation(str(path))
    parts: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        texts.append(" | ".join(cells))
        parts.append(f"## Slide {number}\n\n" + "\n\n".join(texts))
    return "\n\n".join(parts)


CONVERTERS: dict[SourceType, Converter] = {
    "pdf": convert_pdf,
    "md": convert_text,
    "tex": convert_text,
    "ipynb": convert_ipynb,
    "docx": convert_docx,
    "pptx": convert_pptx,
}


@dataclass(frozen=True, slots=True)
class ConversionResult:
    text: str
    status: ConversionStatus
    error: str | None = None


def convert(
    document: RawDocument, converters: dict[SourceType, Converter] = CONVERTERS
) -> ConversionResult:
    """Convert one document; any exception is logged and turned into a failed result."""
    converter = converters.get(document.source_type)
    if converter is None:
        error = f"no converter for source type {document.source_type!r}"
        logger.warning("conversion_failed", source_uri=document.source_uri, error=error)
        return ConversionResult(text="", status="failed", error=error)
    try:
        text = converter(document.path)
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"
        logger.warning("conversion_failed", source_uri=document.source_uri, error=error)
        return ConversionResult(text="", status="failed", error=error)
    # NUL bytes break downstream storage and tokenizers.
    return ConversionResult(text=text.replace("\x00", ""), status="ok")
